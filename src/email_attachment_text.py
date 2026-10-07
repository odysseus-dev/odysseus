"""Bounded local previews of already-authorized email attachments."""
from pathlib import Path


def attachment_text(path, *, max_chars=12000):
    path = Path(path)
    if path.stat().st_size > 20 * 1024 * 1024:
        return {'content_status': 'too_large', 'content_note': 'Attachment exceeds the 20 MB reading limit.'}
    suffix = path.suffix.lower()
    parts = []
    truncated = False
    try:
        if suffix == '.pdf':
            from pypdf import PdfReader
            reader = PdfReader(path)
            if reader.is_encrypted and not reader.decrypt(''):
                return {'content_status': 'encrypted', 'content_note': 'PDF requires a password.'}
            for number, page in enumerate(reader.pages):
                if number >= 50 or sum(map(len, parts)) >= max_chars:
                    truncated = True
                    break
                parts.append(f'Page {number + 1}:\n' + (page.extract_text() or ''))
            if not any(part.split(':\n', 1)[-1].strip() for part in parts):
                return {'content_status': 'needs_ocr', 'content_note': 'No embedded PDF text. Scanned pages require OCR; contents have not been read.'}
        elif suffix in {'.txt', '.md', '.csv', '.tsv', '.json', '.xml', '.log'}:
            with path.open(encoding='utf-8', errors='replace') as file:
                parts.append(file.read(max_chars + 1))
        elif suffix == '.docx':
            from docx import Document
            doc = Document(path)
            parts.extend(p.text for p in doc.paragraphs)
            for table in doc.tables:
                parts.extend('\t'.join(cell.text for cell in row.cells) for row in table.rows)
        elif suffix == '.xlsx':
            from openpyxl import load_workbook
            book = load_workbook(path, read_only=True, data_only=True, keep_links=False)
            try:
                for sheet in book:
                    parts.append(f'Sheet: {sheet.title}')
                    for index, row in enumerate(sheet.iter_rows(values_only=True)):
                        if index >= 1000 or sum(map(len, parts)) >= max_chars:
                            truncated = True
                            break
                        parts.append('\t'.join('' if cell is None else str(cell) for cell in row))
                    if truncated:
                        break
            finally:
                book.close()
        else:
            return {'content_status': 'unsupported', 'content_note': 'This attachment format has no inline text reader.'}
        text = '\n'.join(parts).strip()
        truncated |= len(text) > max_chars
        return {'content': text[:max_chars], 'content_status': 'read' if text else 'empty',
                'content_note': 'Preview truncated; remaining content was not read.' if truncated else ''}
    except Exception as exc:
        return {'content_status': 'failed', 'content_note': f'Attachment text extraction failed ({type(exc).__name__}); contents have not been read.'}
